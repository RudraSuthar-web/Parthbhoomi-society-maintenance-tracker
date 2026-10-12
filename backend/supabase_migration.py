import sys
import openpyxl
import urllib.request
import urllib.error
import json
import os

def load_env(env_path):
    env_vars = {}
    with open(env_path, 'r') as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith('#'):
                key, val = line.split('=', 1)
                env_vars[key.strip()] = val.strip().strip('"').strip("'")
    return env_vars

def parse_excel(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    unit_dues = []
    expenses = []

    # Helper for all months
    all_months = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
    
    def process_sheet(sheet_name, year, start_row, house_col, month_start_col, debit_desc_col, debit_amt_col):
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        
        # Month mapping based on start column
        start_idx = openpyxl.utils.column_index_from_string(month_start_col)
        month_mapping = {}
        for i, month in enumerate(all_months):
            col_letter = openpyxl.utils.get_column_letter(start_idx + i)
            month_mapping[col_letter] = (month, year, f"{year}-{i+1:02d}-01")

        for row in range(start_row, ws.max_row + 1):
            unit_raw = str(ws[f"{house_col}{row}"].value).strip()
            if unit_raw.isdigit():
                unit_num = str(int(unit_raw))
                for col, (month, y, date_cleared) in month_mapping.items():
                    amount_val = str(ws[f"{col}{row}"].value).strip()
                    is_paid = (amount_val == '500' or amount_val == '500.0' or amount_val == '3500' or amount_val == '1500' or amount_val == '4000' or amount_val == '3000' or amount_val == '2000' or (amount_val.isdigit() and int(amount_val) > 0)) 
                    # from 2022 sheet we see amounts like 3500, 1500, 4000 etc. If there is a non-zero positive number, we assume paid.
                    # Wait, let's just say if amount_val is a digit/float and > 0, it's paid.
                    
                    try:
                        amt_float = float(amount_val)
                        is_paid = amt_float > 0
                        amount_to_save = amt_float
                    except ValueError:
                        is_paid = False
                        amount_to_save = 500.0 # Default maintenance

                    unit_dues.append({
                        "tenement_number": unit_num,
                        "month": month,
                        "year": y,
                        "status": "Paid" if is_paid else "Unpaid",
                        "amount": 500.0 if amount_to_save == 0 else amount_to_save, # Standard maintenance is 500
                        "amount_paid": amount_to_save if is_paid else 0,
                        "date_cleared": date_cleared if is_paid else None,
                        "reference": f"EXCEL-IMPORT-{year}" if is_paid else None,
                        "method": "Cash" if is_paid else None,
                    })
            
            # Expenses
            if debit_desc_col and debit_amt_col:
                debit_desc = str(ws[f"{debit_desc_col}{row}"].value or '').strip()
                debit_amt = str(ws[f"{debit_amt_col}{row}"].value or '').strip()
                if debit_desc and debit_amt and debit_desc not in ['DEBIT', 'Total', 'Balance', 'None']:
                    try:
                        expenses.append({
                            "id": f"E-EXCEL-{year}-{row}",
                            "category": "Maintenance",
                            "description": debit_desc,
                            "amount": float(debit_amt),
                            "date": f"{year}-12-31",
                            "drive_link": ""
                        })
                    except ValueError:
                        pass
    
    # 2022: House in A, Starts Row 2, Months start at B, Debit Desc in N, Debit Amt in P
    process_sheet('2022', 2022, 2, 'A', 'B', 'N', 'P')
    # 2023: House in A, Starts Row 4, Months start at B, Debit Desc in N, Debit Amt in Q
    process_sheet('2023', 2023, 4, 'A', 'B', 'N', 'Q')

    # 2025: House in B, Starts Row 1. Months start at C for 5 months only (Aug to Dec).
    # Need custom logic for 2025 and 2026 as previously mapped.
    # 2025 Mapping
    sheet1_mapping = {
        "C": ("August", 2025, "2025-08-01"),
        "D": ("September", 2025, "2025-09-01"),
        "E": ("October", 2025, "2025-10-01"),
        "F": ("November", 2025, "2025-11-01"),
        "G": ("December", 2025, "2025-12-01"),
    }
    if '2025' in wb.sheetnames:
        ws_2025 = wb['2025']
        for row in range(1, ws_2025.max_row + 1):
            unit_raw = str(ws_2025[f"B{row}"].value).strip()
            if unit_raw.isdigit():
                unit_num = str(int(unit_raw))
                for col, (month, year, date_cleared) in sheet1_mapping.items():
                    amount_val = str(ws_2025[f"{col}{row}"].value).strip()
                    try:
                        amt = float(amount_val)
                        is_paid = amt > 0
                    except ValueError:
                        amt = 500.0
                        is_paid = False

                    unit_dues.append({
                        "tenement_number": unit_num,
                        "month": month,
                        "year": year,
                        "status": "Paid" if is_paid else "Unpaid",
                        "amount": 500,
                        "amount_paid": amt if is_paid else 0,
                        "date_cleared": date_cleared if is_paid else None,
                        "reference": "EXCEL-IMPORT-2025" if is_paid else None,
                        "method": "Cash" if is_paid else None,
                    })
            
            debit_desc = str(ws_2025[f"H{row}"].value or '').strip()
            debit_amt = str(ws_2025[f"K{row}"].value or '').strip()
            if debit_desc and debit_amt and debit_desc not in ['DEBIT', 'Total', 'Balance', 'None']:
                try:
                    expenses.append({
                        "id": f"E-EXCEL-2025-{row}",
                        "category": "Maintenance",
                        "description": debit_desc,
                        "amount": float(debit_amt),
                        "date": "2025-12-31",
                        "drive_link": f"https://drive.google.com/file/d/gd-excel-{row}/view"
                    })
                except ValueError:
                    pass

    # 2026 Mapping
    sheet2_mapping = {
        "C": ("January", 2026, "2026-01-01"),
        "D": ("February", 2026, "2026-02-01"),
        "E": ("March", 2026, "2026-03-01"),
        "F": ("April", 2026, "2026-04-01"),
        "G": ("May", 2026, "2026-05-01"),
        "H": ("June", 2026, "2026-06-01"),
        "I": ("July", 2026, "2026-07-01"),
        "J": ("August", 2026, "2026-08-01"),
        "K": ("September", 2026, "2026-09-01"),
        "L": ("October", 2026, "2026-10-01"),
        "M": ("November", 2026, "2026-11-01"),
        "N": ("December", 2026, "2026-12-01"),
    }
    if '2026' in wb.sheetnames:
        ws_2026 = wb['2026']
        for row in range(1, ws_2026.max_row + 1):
            unit_raw = str(ws_2026[f"B{row}"].value).strip()
            if unit_raw.isdigit():
                unit_num = str(int(unit_raw))
                current_month_idx = 6 # July (0-indexed)
                for col, (month, year, date_cleared) in sheet2_mapping.items():
                    amount_val = str(ws_2026[f"{col}{row}"].value).strip()
                    try:
                        amt = float(amount_val)
                        is_paid = amt > 0
                    except ValueError:
                        amt = 500.0
                        is_paid = False

                    m_idx = all_months.index(month)
                    is_future = m_idx > current_month_idx

                    status = "Paid" if is_paid else ("Unbilled" if is_future else "Unpaid")

                    unit_dues.append({
                        "tenement_number": unit_num,
                        "month": month,
                        "year": year,
                        "status": status,
                        "amount": 500,
                        "amount_paid": amt if is_paid else 0,
                        "date_cleared": date_cleared if is_paid else None,
                        "reference": "EXCEL-IMPORT-2026" if is_paid else None,
                        "method": "Cash" if is_paid else None,
                    })

    return unit_dues, expenses

def delete_existing(supabase_url, supabase_key):
    headers = {
        "apikey": supabase_key,
        "Authorization": f"Bearer {supabase_key}",
        "Content-Type": "application/json"
    }

    # Delete 2022, 2023, 2025 and 2026 dues
    for year in [2022, 2023, 2025, 2026]:
        print(f"Deleting existing dues for year {year}...")
        url = f"{supabase_url}/rest/v1/dues?year=eq.{year}"
        req = urllib.request.Request(url, headers=headers, method="DELETE")
        try:
            with urllib.request.urlopen(req) as resp:
                print(f"✓ Deleted dues for {year} (Status: {resp.status})")
        except urllib.error.HTTPError as e:
            print(f"Error deleting dues for {year}: {e.code} {e.reason} - {e.read()}")

        print(f"Deleting existing installments for year {year}...")
        url = f"{supabase_url}/rest/v1/installments?year=eq.{year}"
        req = urllib.request.Request(url, headers=headers, method="DELETE")
        try:
            with urllib.request.urlopen(req) as resp:
                print(f"✓ Deleted installments for {year} (Status: {resp.status})")
        except urllib.error.HTTPError as e:
            print(f"Error deleting installments for {year}: {e.code} {e.reason} - {e.read()}")

    # Delete expenses with EXCEL source
    print(f"Deleting existing Excel-imported expenses...")
    url = f"{supabase_url}/rest/v1/expenses?id=like.E-EXCEL-*"
    req = urllib.request.Request(url, headers=headers, method="DELETE")
    try:
        with urllib.request.urlopen(req) as resp:
            print(f"✓ Deleted Excel expenses (Status: {resp.status})")
    except urllib.error.HTTPError as e:
        print(f"Error deleting expenses: {e.code} {e.reason} - {e.read()}")


def insert_data(supabase_url, supabase_key, unit_dues, expenses):
    headers = {
        "apikey": supabase_key,
        "Authorization": f"Bearer {supabase_key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"
    }
    
    batch_size = 200

    # 1. UPSERT Dues
    print(f"\\nExecuting Bulk UPSERT for {len(unit_dues)} Dues to Supabase REST API...")
    url = f"{supabase_url}/rest/v1/dues?on_conflict=tenement_number,month,year"
    for i in range(0, len(unit_dues), batch_size):
        batch = unit_dues[i:i+batch_size]
        req = urllib.request.Request(url, data=json.dumps(batch).encode("utf-8"), headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=20) as resp:
                print(f"✓ Dues Batch {i//batch_size + 1} Bulk Upsert Status: {resp.status}")
        except urllib.error.HTTPError as e:
            print(f"Supabase Dues Upsert Error: {e.code} {e.reason} - {e.read()}")

    # 2. Insert Installments
    paid_installments = []
    for due in unit_dues:
        if due["status"] == "Paid":
            paid_installments.append({
                "tenement_number": due["tenement_number"],
                "month": due["month"],
                "year": due["year"],
                "amount": due["amount_paid"],
                "date": due["date_cleared"] or f"{due['year']}-12-31",
                "reference": due["reference"],
                "method": "Cash"
            })

    if paid_installments:
        print(f"\\nExecuting Bulk Insert for {len(paid_installments)} Installments...")
        inst_url = f"{supabase_url}/rest/v1/installments"
        for i in range(0, len(paid_installments), batch_size):
            batch = paid_installments[i:i+batch_size]
            inst_req = urllib.request.Request(inst_url, data=json.dumps(batch).encode("utf-8"), headers=headers, method="POST")
            try:
                with urllib.request.urlopen(inst_req, timeout=20) as resp:
                    print(f"✓ Installments Batch {i//batch_size + 1} Bulk Insert Status: {resp.status}")
            except urllib.error.HTTPError as e:
                print(f"Supabase Installments Error: {e.code} {e.reason} - {e.read()}")

    # 3. Insert Expenses
    if expenses:
        print(f"\\nExecuting Bulk UPSERT for {len(expenses)} Expenses...")
        exp_url = f"{supabase_url}/rest/v1/expenses?on_conflict=id"
        exp_req = urllib.request.Request(exp_url, data=json.dumps(expenses).encode("utf-8"), headers=headers, method="POST")
        try:
            with urllib.request.urlopen(exp_req, timeout=20) as resp:
                print(f"✓ Expenses Bulk Upsert Status: {resp.status}")
        except urllib.error.HTTPError as e:
            print(f"Supabase Expenses Upsert Error: {e.code} {e.reason} - {e.read()}")

if __name__ == "__main__":
    excel_file = "/home/hardik/Desktop/Parthbhoomi-society-maintenance-tracker/Parthbhoomi_updated.xlsx"
    env_file = "/home/hardik/Desktop/Parthbhoomi-society-maintenance-tracker/.env"
    
    env_vars = load_env(env_file)
    supabase_url = env_vars.get("SUPABASE_URL")
    supabase_key = env_vars.get("SUPABASE_KEY")
    
    if not supabase_url or not supabase_key:
        print("Error: SUPABASE_URL or SUPABASE_KEY not found in .env")
        sys.exit(1)
    
    print(f"Loading Excel file: {excel_file}")
    dues, expenses = parse_excel(excel_file)
    print(f"Found {len(dues)} dues records and {len(expenses)} expenses records.")
    
    print("\\n--- DELETING OLD ENTRIES ---")
    delete_existing(supabase_url, supabase_key)
    
    print("\\n--- INSERTING NEW ENTRIES ---")
    insert_data(supabase_url, supabase_key, dues, expenses)
    
    print("\\nMigration Complete!")
